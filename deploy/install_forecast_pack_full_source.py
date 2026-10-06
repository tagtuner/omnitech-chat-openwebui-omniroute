#!/usr/bin/env python3
"""Deploy Generate Forecast Pack v2 FULL_FILE_READ + lib + follow_up_gate."""
from __future__ import annotations

import json
import sys
import time
import uuid
from pathlib import Path

import paramiko
from openpyxl import Workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
HOST = "192.168.0.17"
DB = "/var/lib/docker/volumes/open-webui_open-webui-data/_data/webui.db"
# Named user UUID for Forecast Pack ACL (replace with your OWUI user id)
GHAZALA = ""


def _make_smoke_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "AS GB"
    ws.append(["Name", "Subject", "Forecast", "CAIE"])
    ws.append(["Aisha Khan", "Physics", "A", "B"])
    ws.append(["Sara Ali", "Chemistry", "B", "B"])
    ws.append(["Noor", "Mathematics", "A*", "A"])
    ws2 = wb.create_sheet("A2 GB")
    ws2.append(["Name", "Subject", "Forecast", "CAIE"])
    ws2.append(["Zain", "Biology", "A", "A"])
    ws2.append(["Hiba", "Urdu", "C", "B"])
    wb.save(path)


def main() -> int:
    # local unit smoke
    sys.path.insert(0, str(HERE))
    import lib_source_workbook as sw

    smoke_local = HERE / "_tmp_forecast_smoke.xlsx"
    _make_smoke_xlsx(smoke_local)
    book = sw.load_workbook_rows(str(smoke_local))
    built = sw.build_forecast_pack_from_workbook(book, title="Smoke Pack", cohort="GB")
    v = built["verify"]
    assert v["as_rows"] == 3 and v["a2_rows"] == 2 and v["total_rows"] == 5, v
    print("LOCAL_SMOKE_OK", v)
    smoke_local.unlink(missing_ok=True)

    c = paramiko.SSHClient()
    c.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    c.connect(HOST, username="root", timeout=30)

    def run(cmd: str, t: int = 180) -> str:
        print(f"$ {cmd[:160]}")
        _, o, e = c.exec_command(cmd, timeout=t)
        out = (o.read() + e.read()).decode("utf-8", "replace")
        if out.strip():
            print(out[-5000:] if len(out) > 5000 else out)
        return out

    run("mkdir -p /opt/open-webui/tools /opt/open-webui/functions /tmp/owui_smoke")
    sftp = c.open_sftp()
    sftp.put(str(HERE / "lib_source_workbook.py"), "/opt/open-webui/tools/lib_source_workbook.py")
    sftp.put(str(HERE / "generate_forecast_pack.py"), "/opt/open-webui/tools/generate_forecast_pack.py")
    sftp.put(str(HERE / "follow_up_gate.py"), "/opt/open-webui/functions/follow_up_gate.py")
    # smoke workbook on server
    local_smoke = HERE / "_tmp_forecast_smoke_remote.xlsx"
    _make_smoke_xlsx(local_smoke)
    sftp.put(str(local_smoke), "/tmp/owui_smoke/forecast_smoke.xlsx")
    local_smoke.unlink(missing_ok=True)
    sftp.close()

    run("docker cp /opt/open-webui/tools/lib_source_workbook.py open-webui:/app/backend/data/lib_source_workbook.py")
    run("docker cp /opt/open-webui/tools/generate_forecast_pack.py open-webui:/tmp/generate_forecast_pack.py")
    run("docker cp /tmp/owui_smoke/forecast_smoke.xlsx open-webui:/tmp/forecast_smoke.xlsx")

    content = (HERE / "generate_forecast_pack.py").read_text(encoding="utf-8")
    filter_content = (HERE / "follow_up_gate.py").read_text(encoding="utf-8")
    specs = [
        {
            "name": "generate_forecast_pack",
            "description": (
                "Create CAIE forecast Excel pack (Summary+AS+A2+Variance). "
                "When Progress Report spreadsheet attached, MUST use "
                "job=forecast_from_source for COMPLETE disk read (prefer AS GB + A2 GB). "
                "Never invent rows from RAG."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "content": {
                        "type": "string",
                        "description": "JSON only when no source file. Ignored for FULL_FILE_READ.",
                    },
                    "source_file_id": {"type": "string", "description": "Attached file id"},
                    "source_filename": {"type": "string", "description": "Filename hint"},
                    "job": {
                        "type": "string",
                        "description": "forecast_from_source | full_file — FULL_FILE_READ",
                    },
                    "as_sheet": {"type": "string", "description": "Override AS sheet name"},
                    "a2_sheet": {"type": "string", "description": "Override A2 sheet name"},
                    "cohort": {"type": "string", "description": "Default GB"},
                },
                "required": [],
            },
        }
    ]
    payload = {
        "db": DB,
        "ghazala": GHAZALA,
        "content": content,
        "filter_content": filter_content,
        "specs": specs,
    }
    sftp = c.open_sftp()
    with sftp.file("/tmp/forecast_full_payload.json", "w") as f:
        f.write(json.dumps(payload))
    sftp.close()

    runner = r'''
import json, sqlite3, time, uuid
from pathlib import Path
p=json.loads(Path("/tmp/forecast_full_payload.json").read_text(encoding="utf-8"))
con=sqlite3.connect(p["db"]); con.row_factory=sqlite3.Row
now=int(time.time())
admin=con.execute("select id from user where role='admin' limit 1").fetchone()
admin_id=admin["id"] if admin else None
specs=json.dumps(p["specs"])
meta=json.dumps({"description":"Forecast pack FULL_FILE_READ v2","manifest":{"title":"Generate Forecast Pack","author":"OmniTech","version":"2.0.0","requirements":["openpyxl"]}})
row=con.execute("select id from tool where name=?", ("Generate Forecast Pack",)).fetchone()
if row:
    tid=row["id"]
    con.execute("update tool set content=?, specs=?, meta=?, updated_at=? where id=?", (p["content"], specs, meta, now, tid))
    print("TOOL_UPDATED", tid, len(p["content"]))
else:
    tid=str(uuid.uuid4())
    con.execute("insert into tool (id,user_id,name,content,specs,meta,valves,updated_at,created_at) values (?,?,?,?,?,?,?,?,?)",
                (tid, admin_id, "Generate Forecast Pack", p["content"], specs, meta, "{}", now, now))
    print("TOOL_INSERTED", tid)
# Named user only when GHAZALA UUID is set — do NOT grant user:* (locked ACL for Forecast Pack)
for ptype,pid in (("user",p["ghazala"]),):
    if not pid:
        print("SKIP_GRANT no named user id configured")
        continue
    ex=con.execute("select id from access_grant where resource_type='tool' and resource_id=? and principal_type=? and principal_id=? and permission='read'",
                   (tid,ptype,pid)).fetchone()
    if not ex:
        con.execute("insert into access_grant(id,resource_type,resource_id,principal_type,principal_id,permission,created_at) values (?,?,?,?,?,?,?)",
                    (str(uuid.uuid4()),"tool",tid,ptype,pid,"read",now))
        print("GRANT", ptype, pid)
# Drop accidental wildcard if present
n=con.execute("delete from access_grant where resource_type='tool' and resource_id=? and principal_type='user' and principal_id='*'", (tid,)).rowcount
if n:
    print("REVOKED_USER_STAR", n)
for fr in con.execute("select id,name from function where lower(name) like '%follow%'"):
    con.execute("update function set content=?, updated_at=? where id=?", (p["filter_content"], now, fr["id"]))
    print("FILTER_UPDATED", fr["name"])
con.commit(); con.close(); print("DB_OK")
'''
    sftp = c.open_sftp()
    with sftp.file("/tmp/install_forecast_db.py", "w") as f:
        f.write(runner)
    sftp.close()
    run("python3 /tmp/install_forecast_db.py")

    # No full restart required if tool content is from DB on each call — but OWUI caches modules.
    # Light restart safer.
    run("docker restart open-webui")
    for i in range(16):
        time.sleep(3)
        health = run("docker inspect -f '{{.State.Health.Status}}' open-webui").strip()
        api = run(
            "curl -sS -o /dev/null -w '%{http_code}' http://192.168.0.17:3080/api/config"
        ).strip()
        print(f"wait {i}: {health} {api}")
        if "healthy" in health and api.endswith("200"):
            break

    run("docker cp /opt/open-webui/tools/lib_source_workbook.py open-webui:/app/backend/data/lib_source_workbook.py")
    run("docker cp /opt/open-webui/tools/generate_forecast_pack.py open-webui:/tmp/generate_forecast_pack.py")
    run("docker cp /tmp/owui_smoke/forecast_smoke.xlsx open-webui:/tmp/forecast_smoke.xlsx")
    run("bash /opt/open-webui/scripts/apply_name_patch.sh || true")

    smoke = r'''
import asyncio, os, sys, sqlite3, uuid, time, json, hashlib
sys.path.insert(0, "/app/backend/data")
import lib_source_workbook as sw
from importlib.util import spec_from_file_location, module_from_spec

# register smoke file in DB so resolve_file_path works
path = "/tmp/forecast_smoke.xlsx"
data = open(path, "rb").read()
fid = str(uuid.uuid4())
stored = f"{fid}_forecast_smoke.xlsx"
uploads = "/app/backend/data/uploads"
os.makedirs(uploads, exist_ok=True)
open(os.path.join(uploads, stored), "wb").write(data)
db = "/app/backend/data/webui.db"
con = sqlite3.connect(db)
now = int(time.time())
uid = con.execute("select id from user where role='admin' limit 1").fetchone()[0]
meta = {"name": "forecast_smoke.xlsx", "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "size": len(data)}
con.execute(
    "insert into file (id,user_id,filename,meta,created_at,hash,data,updated_at,path) values (?,?,?,?,?,?,?,?,?)",
    (fid, uid, "forecast_smoke.xlsx", json.dumps(meta), now, hashlib.sha256(data).hexdigest(), "{}", now, f"/app/backend/data/uploads/{stored}"),
)
con.commit(); con.close()
print("FILE_ID", fid)

book = sw.load_workbook_rows(f"/app/backend/data/uploads/{stored}")
built = sw.build_forecast_pack_from_workbook(book, title="Smoke Pack", cohort="GB")
print("VERIFY", built["verify"])
assert built["verify"]["total_rows"] == 5

spec = spec_from_file_location("gfp", "/tmp/generate_forecast_pack.py")
mod = module_from_spec(spec); spec.loader.exec_module(mod)
tools = mod.Tools()

async def run():
    out = await tools.generate_forecast_pack(
        content='{"title":"Smoke Forecast Pack"}',
        source_file_id=fid,
        job="forecast_from_source",
        cohort="GB",
        __files__=[{"id": fid, "filename": "forecast_smoke.xlsx", "type": "file"}],
        __user__={"id": uid, "role": "admin"},
    )
    print(out[:1200])
    assert "FULL_FILE_READ" in out or "VERIFY" in out
    assert "5" in out or "AS sheet" in out
    assert "/api/v1/files/" in out
    return out

asyncio.run(run())
print("SMOKE_OK")
'''
    sftp = c.open_sftp()
    with sftp.file("/tmp/smoke_forecast_full.py", "w") as f:
        f.write(smoke)
    sftp.close()
    run("docker cp /tmp/smoke_forecast_full.py open-webui:/tmp/smoke_forecast_full.py")
    run("docker exec open-webui python /tmp/smoke_forecast_full.py")
    c.close()
    print("DONE")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
