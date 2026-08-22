"""
title: Generate Forecast Pack
author: OmniTech
description: Multi-sheet CAIE forecast pack (Summary, AS, A2, Variance) with outcome colours.
requirements: openpyxl
required_open_webui_version: 0.4.0
version: 1.0.0
license: MIT
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Optional

from pydantic import BaseModel, Field

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    _HAS_XLSX = True
except Exception:
    _HAS_XLSX = False

NAVY = "141C2B"
FORECAST_BLUE = "7EA8D9"
CAIE_GOLD = "E8B83A"
WHITE = "FFFFFF"
INK = "1A1A2E"

# Variance = CAIE − Forecast. Ladder A*=7 … U=1.
GRADE = {
    "A*": 7,
    "A": 6,
    "B": 5,
    "C": 4,
    "D": 3,
    "E": 2,
    "U": 1,
}

FILLS = {
    "none": ("D1D5DB", "4B5563"),
    "m2": ("F8D0D8", "9B1C1C"),
    "m1": ("FDE68A", "92400E"),
    "zero": ("C6EFCE", "006100"),
    "p1": ("E2F0D9", "385723"),
    "p2": ("548235", "FFFFFF"),
}

HEADERS = ["Name", "Subject", "Forecast", "CAIE", "Δ", "Outcome"]


def _slugify(text: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "forecast-pack").strip().lower()).strip("-")
    return s[:48] or "forecast-pack"


def _parse_content(content: str | dict) -> dict:
    if isinstance(content, dict):
        return content
    raw = (content or "{}").strip()
    for fence in ("```json", "```"):
        if raw.startswith(fence):
            raw = raw[len(fence) :]
            if raw.endswith("```"):
                raw = raw[:-3]
            raw = raw.strip()
            break
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            return json.loads(m.group(0))
        raise


def _fill(hex_color: str) -> PatternFill:
    h = (hex_color or NAVY).strip().lstrip("#").upper()
    if len(h) != 6:
        h = NAVY
    return PatternFill("solid", fgColor=h)


def _font(bold=False, size=11, color=WHITE, name="Calibri"):
    return Font(bold=bold, size=size, color=color, name=name)


def _thin():
    s = Side(style="thin", color="D1D5DB")
    return Border(left=s, right=s, top=s, bottom=s)


def _norm_grade(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip().upper().replace(" ", "")
    if s in ("A*", "A+", "A STAR", "A-STAR"):
        return "A*"
    if s in GRADE:
        return s
    if s.startswith("A*") or s == "ASTAR":
        return "A*"
    return str(v).strip()


def _grade_n(v: Any) -> Optional[int]:
    g = _norm_grade(v)
    return GRADE.get(g)


def _outcome_from_var(var: Optional[int], has_both: bool) -> str:
    if not has_both or var is None:
        return "No entry made"
    if var == 0:
        return "On target"
    n = abs(var)
    side = "above" if var > 0 else "below"
    if n == 1:
        return f"One grade {side}"
    return f"{n} grades {side}"


def _band(var: Optional[int], has_both: bool) -> str:
    if not has_both or var is None:
        return "none"
    if var <= -2:
        return "m2"
    if var == -1:
        return "m1"
    if var == 0:
        return "zero"
    if var == 1:
        return "p1"
    return "p2"


def _normalize_row(raw: Any) -> dict:
    if isinstance(raw, dict):
        name = raw.get("name") or raw.get("Name") or ""
        subject = raw.get("subject") or raw.get("Subject") or ""
        forecast = _norm_grade(raw.get("forecast") or raw.get("Forecast") or raw.get("F") or "")
        caie = _norm_grade(raw.get("caie") or raw.get("CAIE") or raw.get("C") or "")
        var = raw.get("variance")
        if var in ("", "-", None):
            var = None
        else:
            try:
                var = int(var)
            except (TypeError, ValueError):
                var = None
    elif isinstance(raw, (list, tuple)):
        vals = list(raw) + ["", "", "", "", "", ""]
        name, subject, forecast, caie = vals[0], vals[1], _norm_grade(vals[2]), _norm_grade(vals[3])
        var = None
        try:
            if vals[4] not in ("", "-", None):
                var = int(vals[4])
        except (TypeError, ValueError):
            var = None
    else:
        return {
            "name": str(raw),
            "subject": "",
            "forecast": "",
            "caie": "",
            "variance": "-",
            "outcome": "No entry made",
            "band": "none",
            "var_n": None,
        }

    fn, cn = _grade_n(forecast), _grade_n(caie)
    has_both = fn is not None and cn is not None
    if var is None and has_both:
        var = cn - fn
    outcome = _outcome_from_var(var if has_both else None, has_both)
    return {
        "name": str(name).strip(),
        "subject": str(subject).strip(),
        "forecast": forecast or "",
        "caie": caie or "",
        "variance": "-" if not has_both else var,
        "outcome": outcome,
        "band": _band(var if has_both else None, has_both),
        "var_n": var if has_both else None,
    }


def _on_target_pct(rows: list[dict]) -> tuple[int, int, Optional[float]]:
    scored = [r for r in rows if r.get("band") != "none"]
    if not scored:
        return 0, 0, None
    ot = sum(1 for r in scored if r.get("var_n") == 0)
    return ot, len(scored), round(100.0 * ot / len(scored), 1)


def _aa_pct(rows: list[dict], field: str) -> Optional[float]:
    vals = [r.get(field) for r in rows if r.get(field)]
    if not vals:
        return None
    n = sum(1 for v in vals if v in ("A*", "A"))
    return round(100.0 * n / len(vals), 1)


class Tools:
    def __init__(self):
        self.valves = self.Valves()

    class Valves(BaseModel):
        header_color: str = Field(default=NAVY, description="Header fill hex")

    def _autosize(self, ws, max_width=36):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = 10
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                width = max(width, min(max_width, len(val) + 2))
            ws.column_dimensions[letter].width = width

    def _paint_header(self, ws, row, ncols, fill_hex=NAVY):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = _fill(fill_hex)
            cell.font = _font(bold=True, size=11, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _thin()

    def _apply_band(self, cell, band: str):
        bg, fg = FILLS.get(band, FILLS["none"])
        cell.fill = _fill(bg)
        cell.font = _font(bold=True, size=10, color=fg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _thin()

    def _write_student_sheet(self, ws, title: str, rows: list[dict], header_hex: str):
        ws["A1"] = title
        ws["A1"].font = _font(bold=True, size=16, color=header_hex)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        ot, n, pct = _on_target_pct(rows)
        ws["A2"] = f"Entries scored: {n}  ·  On target: {ot}" + (f" ({pct}%)" if pct is not None else "")
        ws["A2"].font = _font(size=10, color="4B5563")

        hr = 4
        for i, h in enumerate(HEADERS, 1):
            ws.cell(row=hr, column=i, value=h)
        self._paint_header(ws, hr, 6, header_hex)
        # Forecast / CAIE accent on header cells 3–4
        ws.cell(row=hr, column=3).fill = _fill(FORECAST_BLUE)
        ws.cell(row=hr, column=4).fill = _fill(CAIE_GOLD)
        ws.cell(row=hr, column=4).font = _font(bold=True, size=11, color=NAVY)

        alt = _fill("F8FAFC")
        for i, r in enumerate(rows):
            rr = hr + 1 + i
            vals = [r["name"], r["subject"], r["forecast"], r["caie"], r["variance"], r["outcome"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=rr, column=c, value=v)
                cell.border = _thin()
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if i % 2 == 1 and c <= 4:
                    cell.fill = alt
            self._apply_band(ws.cell(row=rr, column=5), r["band"])
            self._apply_band(ws.cell(row=rr, column=6), r["band"])

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A{hr}:F{hr + max(len(rows), 1)}"
        self._autosize(ws)

    def _write_summary(self, ws, spec: dict, as_rows: list[dict], a2_rows: list[dict], header_hex: str):
        title = spec.get("title") or "Forecast Pack"
        ws["A1"] = title
        ws["A1"].font = _font(bold=True, size=20, color=WHITE)
        ws.merge_cells("A1:F1")
        for col in range(1, 7):
            ws.cell(row=1, column=col).fill = _fill(header_hex)
        ws.row_dimensions[1].height = 28

        ws["A2"] = spec.get("subtitle") or "Academic Performance and CAIE Grade Forecast Analysis"
        ws["A2"].font = _font(size=12, color=FORECAST_BLUE)
        ws.merge_cells("A2:F2")

        cohort = str(spec.get("cohort") or "GB").strip() or "GB"
        author = spec.get("author") or "Intellect School"
        day = spec.get("date") or datetime.now(timezone.utc).strftime("%Y-%m-%d")
        ws["A3"] = f"Cohort: {cohort}  ·  {author}  ·  {day}"
        ws["A3"].font = _font(size=10, color="4B5563")

        as_ot, as_n, as_pct = _on_target_pct(as_rows)
        a2_ot, a2_n, a2_pct = _on_target_pct(a2_rows)
        all_rows = as_rows + a2_rows
        ov_ot, ov_n, ov_pct = _on_target_pct(all_rows)

        kpis = spec.get("kpis") if isinstance(spec.get("kpis"), dict) else {}

        def k(key, fallback):
            v = kpis.get(key)
            return fallback if v in (None, "") else v

        rows_kpi = [
            ("KPI", "AS", "A2", "Overall"),
            ("Entries scored", k("as_entries", as_n), k("a2_entries", a2_n), k("overall_entries", ov_n)),
            ("On target count", as_ot, a2_ot, ov_ot),
            ("On target %", k("as_on_target_pct", as_pct), k("a2_on_target_pct", a2_pct), k("overall_on_target_pct", ov_pct)),
            ("A*/A Forecast %", k("as_aa_forecast", _aa_pct(as_rows, "forecast")), k("a2_aa_forecast", _aa_pct(a2_rows, "forecast")), ""),
            ("A*/A CAIE %", k("as_aa_caie", _aa_pct(as_rows, "caie")), k("a2_aa_caie", _aa_pct(a2_rows, "caie")), ""),
        ]
        start = 5
        for r_i, row in enumerate(rows_kpi):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=start + r_i, column=c, value=v)
                cell.border = _thin()
                if r_i == 0:
                    cell.fill = _fill(header_hex)
                    cell.font = _font(bold=True, color=WHITE)
                    cell.alignment = Alignment(horizontal="center")
                elif c == 1:
                    cell.font = _font(bold=True, color=NAVY, size=10)
                else:
                    cell.alignment = Alignment(horizontal="center")
                    if r_i == 3 and c > 1:
                        cell.fill = _fill("C6EFCE")

        ws.cell(row=start, column=2).fill = _fill(FORECAST_BLUE)
        ws.cell(row=start, column=3).fill = _fill(CAIE_GOLD)
        ws.cell(row=start, column=3).font = _font(bold=True, color=NAVY)

        legend_row = start + len(rows_kpi) + 2
        ws.cell(row=legend_row, column=1, value="Outcome colours").font = _font(bold=True, color=NAVY, size=12)
        legend = [
            ("No entry", "none"),
            ("On target", "zero"),
            ("One below", "m1"),
            ("Two+ below", "m2"),
            ("One above", "p1"),
            ("Two+ above", "p2"),
        ]
        for i, (lab, band) in enumerate(legend):
            cell = ws.cell(row=legend_row + 1, column=1 + i, value=lab)
            self._apply_band(cell, band)

        note_row = legend_row + 3
        ws.cell(row=note_row, column=1, value="Sheets").font = _font(bold=True, color=NAVY, size=12)
        ws.cell(row=note_row + 1, column=1, value="Summary · AS · A2 · Variance (intervention list)")
        ws.cell(row=note_row + 2, column=1, value="Δ = CAIE − Forecast  (A*=7 … U=1). Empty CAIE or forecast → No entry made.")
        ws.cell(row=note_row + 2, column=1).font = _font(size=9, color="6B7280")

        self._autosize(ws)
        ws.column_dimensions["A"].width = 28

    def _write_variance(self, ws, as_rows: list[dict], a2_rows: list[dict], actions: list, header_hex: str):
        ws["A1"] = "Variance — intervention list"
        ws["A1"].font = _font(bold=True, size=16, color=header_hex)
        ws.merge_cells("A1:G1")
        ws["A2"] = "Not on target (and has both grades). Worst Δ first."
        ws["A2"].font = _font(size=10, color="4B5563")

        headers = ["Level", *HEADERS]
        hr = 4
        for i, h in enumerate(headers, 1):
            ws.cell(row=hr, column=i, value=h)
        self._paint_header(ws, hr, 7, header_hex)
        ws.cell(row=hr, column=4).fill = _fill(FORECAST_BLUE)
        ws.cell(row=hr, column=5).fill = _fill(CAIE_GOLD)
        ws.cell(row=hr, column=5).font = _font(bold=True, color=NAVY)

        flagged = []
        for level, rows in (("AS", as_rows), ("A2", a2_rows)):
            for r in rows:
                if r["band"] not in ("none", "zero"):
                    flagged.append((level, r))
        flagged.sort(key=lambda x: (x[1]["var_n"] is None, x[1]["var_n"] if x[1]["var_n"] is not None else 0, x[0], x[1]["subject"], x[1]["name"]))

        alt = _fill("F8FAFC")
        for i, (level, r) in enumerate(flagged):
            rr = hr + 1 + i
            vals = [level, r["name"], r["subject"], r["forecast"], r["caie"], r["variance"], r["outcome"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=rr, column=c, value=v)
                cell.border = _thin()
                if i % 2 == 1 and c <= 5:
                    cell.fill = alt
            self._apply_band(ws.cell(row=rr, column=6), r["band"])
            self._apply_band(ws.cell(row=rr, column=7), r["band"])

        last = hr + 1 + max(len(flagged), 1)
        act_row = last + 2
        ws.cell(row=act_row, column=1, value="Actions / recommendations").font = _font(bold=True, size=12, color=header_hex)
        if not actions:
            actions = ["Review Variance sheet with heads of subject.", "Priority: two or more grades below target."]
        for i, a in enumerate(actions):
            ws.cell(row=act_row + 1 + i, column=1, value=f"• {a}")
            ws.merge_cells(start_row=act_row + 1 + i, start_column=1, end_row=act_row + 1 + i, end_column=7)

        ws.freeze_panes = "A5"
        self._autosize(ws)
        ws.column_dimensions["B"].width = 22

    def _build_xlsx(self, spec: dict) -> bytes:
        header_hex = (spec.get("header_color") or self.valves.header_color or NAVY).lstrip("#")
        as_rows = [_normalize_row(r) for r in (spec.get("as_rows") or spec.get("AS") or [])]
        a2_rows = [_normalize_row(r) for r in (spec.get("a2_rows") or spec.get("A2") or [])]
        actions = spec.get("actions") or spec.get("notes") or spec.get("takeaways") or []
        if isinstance(actions, str):
            actions = [actions]

        wb = Workbook()
        ws_sum = wb.active
        ws_sum.title = "Summary"
        self._write_summary(ws_sum, spec, as_rows, a2_rows, header_hex)

        ws_as = wb.create_sheet("AS")
        self._write_student_sheet(ws_as, "AS", as_rows, header_hex)
        ws_a2 = wb.create_sheet("A2")
        self._write_student_sheet(ws_a2, "A2", a2_rows, header_hex)
        ws_var = wb.create_sheet("Variance")
        self._write_variance(ws_var, as_rows, a2_rows, actions, header_hex)

        meta = wb.create_sheet("_meta")
        meta["A1"] = "Pack"
        meta["B1"] = "forecast_v1"
        meta["A2"] = "Title"
        meta["B2"] = spec.get("title") or ""
        meta["A3"] = "Cohort"
        meta["B3"] = spec.get("cohort") or "GB"
        meta["A4"] = "Generated"
        meta["B4"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        meta.sheet_state = "hidden"

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _save(self, data: bytes, *, title: str, user_dict: Optional[dict]):
        slug = _slugify(title)
        day = datetime.now(timezone.utc).strftime("%Y%m%d")
        short = uuid.uuid4().hex[:6]
        filename = f"forecast-pack-{slug}_{day}_{short}.xlsx"
        fid = str(uuid.uuid4())
        stored = f"{fid}_{filename}"
        uploads = "/app/backend/data/uploads"
        cache = "/app/backend/data/cache/files"
        try:
            os.makedirs(uploads, mode=0o775, exist_ok=True)
            os.makedirs(cache, mode=0o775, exist_ok=True)
            with open(os.path.join(uploads, stored), "wb") as fh:
                fh.write(data)
            with open(os.path.join(cache, filename), "wb") as fh:
                fh.write(data)
            uid = user_dict.get("id") if isinstance(user_dict, dict) else None
            db_path = "/app/backend/data/webui.db"
            file_hash = hashlib.sha256(data).hexdigest()
            now = int(datetime.now(timezone.utc).timestamp())
            meta = {
                "name": filename,
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "size": len(data),
                "file_hash": file_hash,
                "data": {},
            }
            con = sqlite3.connect(db_path)
            try:
                if not uid:
                    row = con.execute("select id from user where role='admin' limit 1").fetchone()
                    uid = row[0] if row else "system"
                cols = [r[1] for r in con.execute("pragma table_info(file)")]
                vals = {
                    "id": fid,
                    "user_id": uid,
                    "filename": filename,
                    "meta": json.dumps(meta),
                    "created_at": now,
                    "hash": file_hash,
                    "data": json.dumps({}),
                    "updated_at": now,
                    "path": f"/app/backend/data/uploads/{stored}",
                }
                use = [k for k in vals if k in cols]
                con.execute(
                    f"insert into file ({','.join(use)}) values ({','.join('?' for _ in use)})",
                    [vals[k] for k in use],
                )
                con.commit()
            finally:
                con.close()
            return filename, f"/api/v1/files/{fid}/content", None
        except Exception as exc:
            return filename, None, str(exc)

    async def generate_forecast_pack(
        self,
        content: str = "{}",
        __event_emitter__: Any = None,
        __user__: Optional[dict] = None,
        __request__: Any = None,
    ) -> str:
        """Create a branded multi-sheet CAIE forecast Excel pack and return a download link.

        Use when the user wants a forecast pack, progress pack, AS/A2 CAIE workbook,
        or Students Progress Report style Excel with Summary + AS + A2 + Variance.

        Call this AFTER the one clarification round (or immediately if they said
        use defaults / already answered). Defaults: cohort GB, 4 sheets, colours on.

        `content` MUST be a SINGLE JSON string (no markdown fence):
        {
          "title": "Students Progress Report 2026",
          "subtitle": "optional",
          "author": "Intellect School",
          "cohort": "GB",
          "kpis": {"as_on_target_pct": 43, "a2_on_target_pct": 58, "overall_on_target_pct": 51},
          "as_rows": [{"name":"...","subject":"...","forecast":"A","caie":"B"}],
          "a2_rows": [{"name":"...","subject":"...","forecast":"A","caie":"A"}],
          "actions": ["Computer Science revision sessions", "..."]
        }
        Variance and Outcome are computed if omitted (Δ = CAIE − Forecast, A*=7).
        Reproduce the markdown download link EXACTLY (do NOT prefix sandbox:).
        """
        if not _HAS_XLSX:
            return (
                "[TOOL_RESULT — use as final reply]\n\n"
                "I couldn't generate the forecast pack: openpyxl is not installed."
            )
        try:
            spec = _parse_content(content)
        except Exception as exc:
            return f"[TOOL_RESULT — use as final reply]\n\nInvalid JSON for generate_forecast_pack: {exc}"

        if __event_emitter__:
            try:
                await __event_emitter__(
                    {"type": "status", "data": {"description": "Building forecast pack…", "done": False}}
                )
            except Exception:
                pass

        try:
            data = self._build_xlsx(spec)
        except Exception as exc:
            return f"[TOOL_RESULT — use as final reply]\n\nForecast pack render failed: {exc}"

        fname, url, err = self._save(
            data, title=str(spec.get("title") or "forecast-pack"), user_dict=__user__
        )
        if err or not url:
            return f"[TOOL_RESULT — use as final reply]\n\nPack built but save failed: {err}"

        kb = max(1, len(data) // 1024)
        if __event_emitter__:
            try:
                await __event_emitter__(
                    {
                        "type": "message",
                        "data": {
                            "content": (
                                f"\n\n---\n\n📊 **Forecast pack ready** · {kb} KB\n\n"
                                f"⬇️ [Download {fname}]({url})\n\n---\n"
                            )
                        },
                    }
                )
                await __event_emitter__(
                    {"type": "status", "data": {"description": "Forecast pack ready", "done": True}}
                )
            except Exception:
                pass

        return (
            "[TOOL_RESULT — reproduce the markdown link EXACTLY as written below "
            "(do NOT change the URL, do NOT prefix sandbox:). Do not include this line.]\n\n"
            "Here is the forecast pack:\n\n"
            f"[{fname}]({url})"
        )
