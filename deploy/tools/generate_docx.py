"""
title: Generate Word
author: OmniTech
description: Generate native Word (.docx). Prefer FULL attached-file read via source_file_id/job.
requirements: python-docx
required_open_webui_version: 0.4.0
version: 2.0.0
license: MIT
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import sys
import uuid
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Optional

from pydantic import BaseModel, Field

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt, RGBColor, Cm

    _HAS_DOCX = True
except Exception:
    _HAS_DOCX = False


def _load_source_lib():
    for p in (
        "/app/backend/data",
        "/opt/open-webui/tools",
        os.path.dirname(os.path.abspath(__file__)),
    ):
        if p and p not in sys.path and os.path.isdir(p):
            sys.path.insert(0, p)
    import lib_source_workbook as sw  # type: ignore

    return sw


def _slugify(text: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "document").strip().lower()).strip("-")
    return (s[:48] or "document")


def _parse_content(content: str | dict) -> dict:
    if isinstance(content, dict):
        return content
    raw = (content or "{}").strip()
    for fence in ("```json", "```"):
        if raw.startswith(fence):
            raw = raw[len(fence) :]
            if raw.endswith("```"):
                raw = raw[:-3]
            raw = raw.strip()
            break
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            return json.loads(m.group(0))
        raise


def _rgb(hex_color: str, default=(20, 28, 43)):
    h = (hex_color or "").strip().lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    if len(h) != 6:
        return RGBColor(*default)
    try:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except ValueError:
        return RGBColor(*default)


def _set_cell_shading(cell, hex_color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#").upper())
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


class Tools:
    def __init__(self):
        self.valves = self.Valves()

    class Valves(BaseModel):
        primary_color: str = Field(default="141C2B", description="Heading color hex")
        accent_color: str = Field(default="7EA8D9", description="Accent / table header hex")

    def _build_docx(self, spec: dict) -> bytes:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

        primary = _rgb(spec.get("primary") or self.valves.primary_color)
        accent_hex = (spec.get("accent") or self.valves.accent_color).lstrip("#")

        title = str(spec.get("title") or "Document")
        p = doc.add_heading(title, level=0)
        for run in p.runs:
            run.font.color.rgb = primary

        if spec.get("subtitle"):
            sub = doc.add_paragraph(str(spec["subtitle"]))
            sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in sub.runs:
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

        meta = []
        if spec.get("author"):
            meta.append(str(spec["author"]))
        meta.append(str(spec.get("date") or datetime.now(timezone.utc).strftime("%Y-%m-%d")))
        mp = doc.add_paragraph(" · ".join(meta))
        for run in mp.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        doc.add_paragraph("")

        sections = spec.get("sections") or spec.get("pages") or []
        if not sections and spec.get("slides"):
            for s in spec["slides"]:
                sections.append(
                    {
                        "heading": s.get("title") or s.get("heading") or "Section",
                        "body": s.get("body") or "",
                        "bullets": s.get("bullets") or s.get("points") or [],
                        "table": s.get("table"),
                    }
                )

        for sec in sections:
            if not isinstance(sec, dict):
                continue
            heading = sec.get("heading") or sec.get("title")
            if heading:
                h = doc.add_heading(str(heading), level=1)
                for run in h.runs:
                    run.font.color.rgb = primary
            if sec.get("subheading"):
                h2 = doc.add_heading(str(sec["subheading"]), level=2)
                for run in h2.runs:
                    run.font.color.rgb = primary

            body = sec.get("body") or sec.get("text") or sec.get("content") or ""
            if body:
                for para in str(body).split("\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())

            bullets = sec.get("bullets") or sec.get("points") or sec.get("items") or []
            for b in bullets:
                if isinstance(b, dict):
                    line = b.get("text") or b.get("title") or b.get("description") or str(b)
                else:
                    line = str(b)
                doc.add_paragraph(line, style="List Bullet")

            # KPI line
            stats = sec.get("stats") or sec.get("kpis") or []
            if stats:
                bits = []
                for st in stats:
                    if isinstance(st, dict):
                        bits.append(f"{st.get('label', '')}: {st.get('value', '')}".strip(": "))
                    else:
                        bits.append(str(st))
                kp = doc.add_paragraph(" | ".join(bits))
                for run in kp.runs:
                    run.bold = True
                    run.font.size = Pt(11)

            table = sec.get("table")
            if isinstance(table, dict):
                headers = table.get("headers") or table.get("columns") or []
                rows = table.get("rows") or []
                if headers:
                    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    tbl.style = "Table Grid"
                    hdr_cells = tbl.rows[0].cells
                    for i, htxt in enumerate(headers):
                        hdr_cells[i].text = str(htxt)
                        for paragraph in hdr_cells[i].paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                run.font.size = Pt(9)
                        try:
                            _set_cell_shading(hdr_cells[i], accent_hex if len(accent_hex) == 6 else "141C2B")
                        except Exception:
                            pass
                    for r_i, row in enumerate(rows):
                        cells = tbl.rows[r_i + 1].cells
                        if isinstance(row, dict):
                            vals = [row.get(h, "") for h in headers]
                        else:
                            vals = list(row) if isinstance(row, (list, tuple)) else [row]
                        while len(vals) < len(headers):
                            vals.append("")
                        for c_i, val in enumerate(vals[: len(headers)]):
                            cells[c_i].text = str(val)
                            for paragraph in cells[c_i].paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(9)
                    doc.add_paragraph("")

        closing = spec.get("closing")
        if isinstance(closing, dict):
            doc.add_heading(str(closing.get("title") or "Recommendations"), level=1)
            for b in closing.get("takeaways") or closing.get("bullets") or []:
                doc.add_paragraph(str(b), style="List Bullet")
            if closing.get("contact"):
                doc.add_paragraph(str(closing["contact"]))

        if not sections and not closing:
            doc.add_paragraph("No sections provided in the JSON spec.")

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def _save(self, data: bytes, *, title: str, user_dict: Optional[dict]):
        slug = _slugify(title)
        day = datetime.now(timezone.utc).strftime("%Y%m%d")
        short = uuid.uuid4().hex[:6]
        filename = f"document-{slug}_{day}_{short}.docx"
        fid = str(uuid.uuid4())
        stored = f"{fid}_{filename}"
        uploads = "/app/backend/data/uploads"
        cache = "/app/backend/data/cache/files"
        try:
            os.makedirs(uploads, mode=0o775, exist_ok=True)
            os.makedirs(cache, mode=0o775, exist_ok=True)
            with open(os.path.join(uploads, stored), "wb") as fh:
                fh.write(data)
            with open(os.path.join(cache, filename), "wb") as fh:
                fh.write(data)
            uid = user_dict.get("id") if isinstance(user_dict, dict) else None
            db_path = "/app/backend/data/webui.db"
            file_hash = hashlib.sha256(data).hexdigest()
            now = int(datetime.now(timezone.utc).timestamp())
            meta = {
                "name": filename,
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "size": len(data),
                "file_hash": file_hash,
                "data": {},
            }
            con = sqlite3.connect(db_path)
            try:
                if not uid:
                    row = con.execute("select id from user where role='admin' limit 1").fetchone()
                    uid = row[0] if row else "system"
                cols = [r[1] for r in con.execute("pragma table_info(file)")]
                vals = {
                    "id": fid,
                    "user_id": uid,
                    "filename": filename,
                    "meta": json.dumps(meta),
                    "created_at": now,
                    "hash": file_hash,
                    "data": json.dumps({}),
                    "updated_at": now,
                    "path": f"/app/backend/data/uploads/{stored}",
                }
                use = [k for k in vals if k in cols]
                con.execute(
                    f"insert into file ({','.join(use)}) values ({','.join('?' for _ in use)})",
                    [vals[k] for k in use],
                )
                con.commit()
            finally:
                con.close()
            return filename, f"/api/v1/files/{fid}/content", None
        except Exception as exc:
            return filename, None, str(exc)

    async def generate_docx(
        self,
        content: str = "{}",
        source_file_id: str = "",
        source_filename: str = "",
        job: str = "",
        sheet_name: str = "",
        group_by: str = "Subject",
        __files__: Optional[list] = None,
        __event_emitter__: Any = None,
        __user__: Optional[dict] = None,
        __request__: Any = None,
    ) -> str:
        """Create a native Word (.docx) document and return a download link.

        FULL-FILE RULE: when a spreadsheet is attached, call with
        job='subject_group_report' (reads COMPLETE file on disk). Never invent
        tables from RAG snippets.

        Optional blank-doc JSON content when no source job:
        { "title":"...", "sections":[{"heading":"...","bullets":["..."]}] }
        """
        if not _HAS_DOCX:
            return (
                "[TOOL_RESULT — use as final reply]\n\n"
                "I couldn't generate the Word file: python-docx is not installed."
            )

        verify_md = ""
        job_l = (job or "").strip().lower()
        if (source_file_id or source_filename or job_l) and job_l in (
            "",
            "subject_group_report",
            "subjectwise",
            "subject_wise",
            "subject-wise",
        ):
            job_l = job_l or "subject_group_report"

        if job_l in ("subject_group_report", "subjectwise", "subject_wise", "subject-wise"):
            try:
                sw = _load_source_lib()
                uid = __user__.get("id") if isinstance(__user__, dict) else None
                fid, filename, path = sw.resolve_file_path(
                    source_file_id=source_file_id,
                    source_filename=source_filename,
                    files=__files__ or [],
                    user_id=uid,
                )
                book = sw.load_workbook_rows(path)
                title = "Subject-wise Report"
                try:
                    meta = _parse_content(content) if content and content.strip() not in ("{}", "") else {}
                    if isinstance(meta, dict) and meta.get("title"):
                        title = str(meta["title"])
                except Exception:
                    pass
                built = sw.build_subject_group_report(
                    book,
                    sheet_name=sheet_name or "Form Responses 1",
                    group_by=group_by or "Subject",
                    title=title,
                )
                spec = built["docx"]
                verify_md = sw.verify_block_markdown(built["verify"], fid, filename)
            except Exception as exc:
                return (
                    "[TOOL_RESULT — use as final reply]\n\n"
                    f"FULL_FILE_READ failed: {exc}"
                )
        else:
            try:
                spec = _parse_content(content)
            except Exception as exc:
                return f"[TOOL_RESULT — use as final reply]\n\nInvalid JSON for generate_docx: {exc}"
            if __files__:
                for item in __files__ or []:
                    if not isinstance(item, dict):
                        continue
                    n = str(item.get("filename") or item.get("name") or "").lower()
                    if n.endswith((".xlsx", ".xlsm", ".xls", ".csv")):
                        return (
                            "[TOOL_RESULT — use as final reply]\n\n"
                            "Spreadsheet attached — call generate_docx with "
                            "job='subject_group_report' for FULL_FILE_READ."
                        )

        if __event_emitter__:
            try:
                await __event_emitter__(
                    {"type": "status", "data": {"description": "Building Word document…", "done": False}}
                )
            except Exception:
                pass

        try:
            data = self._build_docx(spec)
        except Exception as exc:
            return f"[TOOL_RESULT — use as final reply]\n\nDOCX render failed: {exc}"

        fname, url, err = self._save(data, title=str(spec.get("title") or "document"), user_dict=__user__)
        if err or not url:
            return f"[TOOL_RESULT — use as final reply]\n\nDOCX built but save failed: {err}"

        kb = max(1, len(data) // 1024)
        head = verify_md + "\n\n" if verify_md else ""
        if __event_emitter__:
            try:
                await __event_emitter__(
                    {
                        "type": "message",
                        "data": {
                            "content": (
                                f"\n\n---\n\n📝 **Word document ready** · {kb} KB\n\n"
                                f"{head}"
                                f"⬇️ [Download {fname}]({url})\n\n---\n"
                            )
                        },
                    }
                )
                await __event_emitter__(
                    {"type": "status", "data": {"description": "DOCX ready", "done": True}}
                )
            except Exception:
                pass

        return (
            "[TOOL_RESULT — reproduce the markdown link EXACTLY as written below "
            "(do NOT change the URL, do NOT prefix sandbox:). Do not include this line.]\n\n"
            f"{head}"
            "Here is the Word document:\n\n"
            f"[{fname}]({url})"
        )
